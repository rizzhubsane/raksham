"""Deck text extraction."""

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from deckeval.content import collect_shape_text, extract_slides, normalize_url, pair_images


class _Paragraph:
    def __init__(self, text: str) -> None:
        self.text = text


class _Frame:
    def __init__(self, text: str) -> None:
        self.paragraphs = [_Paragraph(text)]


class _Cell:
    def __init__(self, text: str) -> None:
        self.text = text


class _Row:
    def __init__(self, cells: list[_Cell]) -> None:
        self.cells = cells


class _Table:
    def __init__(self, rows: list[_Row]) -> None:
        self.rows = rows


class _Shape:
    def __init__(self, shape_type, text=None, children=None, table=None):
        self.shape_type = shape_type
        self.shapes = children or []
        self.has_table = table is not None
        self.table = table
        self.has_text_frame = text is not None
        self.has_chart = False
        self.text_frame = _Frame(text or "")


def test_grouped_shapes_and_tables_are_extracted():
    child = _Shape(MSO_SHAPE_TYPE.TEXT_BOX, text="Inside the group")
    group = _Shape(MSO_SHAPE_TYPE.GROUP, children=[child])
    assert collect_shape_text(group) == ["Inside the group"]

    table = _Table([_Row([_Cell("Cost"), _Cell("High")])])
    assert collect_shape_text(_Shape(MSO_SHAPE_TYPE.TABLE, table=table)) == ["Cost | High"]
    assert collect_shape_text(_Shape(MSO_SHAPE_TYPE.TEXT_BOX, text="Click to add text")) == []


def test_extract_slides_reads_title_body_table_and_notes(tmp_path: Path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "Problem"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
    box.text_frame.text = "Hospitals lose time"
    table_shape = slide.shapes.add_table(1, 2, Inches(1), Inches(3.2), Inches(4), Inches(0.6))
    table_shape.table.cell(0, 0).text = "Cost"
    table_shape.table.cell(0, 1).text = "High"
    slide.notes_slide.notes_text_frame.text = "Emphasize the delay"
    path = tmp_path / "deck.pptx"
    presentation.save(path)

    slides = extract_slides(path)
    assert len(slides) == 1
    assert slides[0].title == "Problem"
    assert "Hospitals lose time" in slides[0].body
    assert "Cost | High" in slides[0].body
    assert slides[0].notes == "Emphasize the delay"


def test_normalize_url_and_image_pairing():
    assert normalize_url("-") is None
    assert normalize_url("  ") is None
    assert normalize_url("notaurl") is None
    assert normalize_url(" https://cdn.example.com/a.pptx ") == "https://cdn.example.com/a.pptx"

    assert pair_images(0, []) == ([], False)
    missing, text_only = pair_images(2, [])
    assert missing == [] and text_only is True
    paired, text_only = pair_images(1, [(Path("slide.png"), "png")])
    assert text_only is False
    assert paired[0].number == 1
    assert paired[0].format == "png"
