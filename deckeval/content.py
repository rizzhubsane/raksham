"""Download decks, extract text, and render slide images."""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import signal
import subprocess
import tempfile
import threading
import time
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import quote

import requests
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

# One try plus 3 retries.
DOWNLOAD_ATTEMPTS = 4
DOWNLOAD_TIMEOUT = (10, 90)
MAX_DOWNLOAD_BYTES = 100 * 1024 * 1024
MAX_SLIDE_CHARS = 8_000
MAX_DECK_CHARS = 80_000
_CACHE_VERSION = 1
_MISSING_URLS = {"", "-", "—", "n/a", "na", "none", "null"}
_PLACEHOLDERS = {
    "click to add title",
    "click to add text",
    "click to add subtitle",
    "click to add notes",
}
_SOFFICE_CANDIDATES = (
    "soffice",
    "libreoffice",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/usr/bin/soffice",
    "/usr/bin/libreoffice",
)
_RENDER_LOCK = threading.Lock()
_URL_LOCKS: dict[str, threading.Lock] = {}
_URL_LOCKS_GUARD = threading.Lock()


class DownloadError(Exception):
    """The deck could not be downloaded or opened."""

    def __init__(self, reason: str, *, retryable: bool = False) -> None:
        super().__init__(reason)
        self.reason = reason
        self.retryable = retryable


class RenderError(Exception):
    """LibreOffice or poppler could not produce slide images."""


@dataclass(frozen=True)
class SlideText:
    number: int
    title: str
    body: str
    notes: str


@dataclass(frozen=True)
class SlideImage:
    number: int
    path: Path
    format: str


@dataclass(frozen=True)
class DeckExtract:
    slides: tuple[SlideText, ...]
    images: tuple[SlideImage, ...]
    text_only: bool


def normalize_url(value: str | None) -> str | None:
    """Return an http(s) URL, or None when the cell is blank or not a link."""
    if value is None:
        return None
    text = value.strip()
    if text.lower() in _MISSING_URLS:
        return None
    if not text.lower().startswith(("http://", "https://")):
        return None
    return text


def find_soffice() -> str | None:
    for candidate in _SOFFICE_CANDIDATES:
        if os.path.isabs(candidate):
            if os.path.isfile(candidate) and os.access(candidate, os.X_OK):
                return candidate
            continue
        found = shutil.which(candidate)
        if found:
            return found
    return None


def poppler_available() -> bool:
    return shutil.which("pdftoppm") is not None or shutil.which("pdfinfo") is not None


def deck_to_text(slides: tuple[SlideText, ...] | list[SlideText]) -> str:
    parts: list[str] = []
    for slide in slides:
        body = slide.body[:MAX_SLIDE_CHARS]
        notes = slide.notes[:MAX_SLIDE_CHARS]
        parts.append(
            f"SLIDE {slide.number}\nTitle: {slide.title}\nText:\n{body}\nSpeaker notes:\n{notes}"
        )
    text = "\n\n".join(parts)
    if len(text) > MAX_DECK_CHARS:
        return text[:MAX_DECK_CHARS] + "\n[truncated]"
    return text


def collect_shape_text(shape: object) -> list[str]:
    """Collect text from a shape, including tables and grouped shapes."""
    texts: list[str] = []
    shape_type = getattr(shape, "shape_type", None)
    if shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in getattr(shape, "shapes", []):
            texts.extend(collect_shape_text(child))
        return texts
    if getattr(shape, "has_table", False):
        table = shape.table
        for row in table.rows:
            cells = [_clean(cell.text) for cell in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            if line:
                texts.append(line)
        return texts
    if getattr(shape, "has_chart", False):
        chart_title = _chart_title(shape)
        if chart_title:
            texts.append(chart_title)
    if getattr(shape, "has_text_frame", False):
        frame = shape.text_frame
        for paragraph in frame.paragraphs:
            line = _clean(paragraph.text)
            if line:
                texts.append(line)
    return texts


def extract_slides(pptx_path: Path) -> tuple[SlideText, ...]:
    """Read titles, body text, tables, grouped shapes, and speaker notes."""
    try:
        presentation = Presentation(str(pptx_path))
    except Exception as exc:
        raise DownloadError("invalid_pptx") from exc
    slides: list[SlideText] = []
    for number, slide in enumerate(presentation.slides, start=1):
        title, title_id = _slide_title(slide)
        body_lines: list[str] = []
        for shape in slide.shapes:
            if title_id is not None and getattr(shape, "shape_id", None) == title_id:
                continue
            body_lines.extend(collect_shape_text(shape))
        slides.append(
            SlideText(
                number=number,
                title=title,
                body="\n".join(body_lines),
                notes=_slide_notes(slide),
            )
        )
    return tuple(slides)


def pair_images(
    slide_count: int, rendered: list[tuple[Path, str]]
) -> tuple[list[SlideImage], bool]:
    """Pair rendered pages with slides. A count mismatch falls back to text only."""
    if slide_count == 0 and not rendered:
        return [], False
    if len(rendered) != slide_count or slide_count == 0:
        return [], True
    images = [
        SlideImage(number=index, path=path, format=image_format)
        for index, (path, image_format) in enumerate(rendered, start=1)
    ]
    return images, False


def load_deck(url: str, cache_dir: Path) -> DeckExtract:
    """Download and extract a deck, reusing `.cache/` entries for the same URL."""
    key = hashlib.sha256(url.strip().encode("utf-8")).hexdigest()[:32]
    deck_dir = cache_dir / key
    with _lock_for(key):
        deck_dir.mkdir(parents=True, exist_ok=True)
        pptx_path = deck_dir / "deck.pptx"
        if not pptx_path.is_file():
            download_pptx(url, pptx_path)
        try:
            slides = _load_or_extract_slides(deck_dir, pptx_path)
        except DownloadError:
            pptx_path.unlink(missing_ok=True)
            raise
        cached_images = _load_cached_images(deck_dir, len(slides))
        if cached_images is not None:
            return DeckExtract(slides=slides, images=tuple(cached_images), text_only=False)
        try:
            rendered = render_slides(pptx_path, deck_dir / "render")
            images, text_only = pair_images(len(slides), rendered)
        except RenderError:
            images, text_only = [], True
        if not text_only:
            _write_manifest(deck_dir, images)
        return DeckExtract(slides=slides, images=tuple(images), text_only=text_only)


def download_pptx(url: str, dest: Path) -> None:
    """Download a .pptx with a timeout and up to 3 retries."""
    last_error: DownloadError | None = None
    for attempt in range(1, DOWNLOAD_ATTEMPTS + 1):
        try:
            _download_once(url, dest)
            return
        except DownloadError as exc:
            last_error = exc
            if attempt == DOWNLOAD_ATTEMPTS or not exc.retryable:
                raise
            time.sleep(2 ** (attempt - 1))
    raise last_error or DownloadError("request_failed")


def render_slides(pptx_path: Path, output_dir: Path, timeout: int = 180) -> list[tuple[Path, str]]:
    """Convert pptx to PDF with LibreOffice, then PDF pages to images."""
    binary = find_soffice()
    if binary is None:
        raise RenderError("libreoffice_not_installed")
    if not poppler_available():
        raise RenderError("poppler_not_installed")
    output_dir.mkdir(parents=True, exist_ok=True)
    with _RENDER_LOCK:
        with tempfile.TemporaryDirectory(prefix="lo-profile-") as profile:
            _convert_to_pdf(binary, pptx_path, output_dir, Path(profile), timeout)
        pdf_path = _find_pdf(output_dir, pptx_path)
        return _pdf_to_images(pdf_path, output_dir / "slides")


def _download_once(url: str, dest: Path) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    partial = dest.with_suffix(".part")
    try:
        with requests.get(
            url,
            timeout=DOWNLOAD_TIMEOUT,
            stream=True,
            headers={"User-Agent": "eval-decks/1.0"},
        ) as response:
            try:
                response.raise_for_status()
            except requests.HTTPError as exc:
                status = exc.response.status_code if exc.response is not None else 0
                retryable = status in {408, 429} or status >= 500
                raise DownloadError(f"http_{status}", retryable=retryable) from exc
            size = 0
            with partial.open("wb") as handle:
                for chunk in response.iter_content(chunk_size=256 * 1024):
                    if not chunk:
                        continue
                    size += len(chunk)
                    if size > MAX_DOWNLOAD_BYTES:
                        raise DownloadError("too_large")
                    handle.write(chunk)
    except DownloadError:
        partial.unlink(missing_ok=True)
        raise
    except requests.Timeout as exc:
        partial.unlink(missing_ok=True)
        raise DownloadError("timeout", retryable=True) from exc
    except requests.RequestException as exc:
        partial.unlink(missing_ok=True)
        raise DownloadError("request_failed", retryable=True) from exc

    if not _is_zip(partial):
        partial.unlink(missing_ok=True)
        raise DownloadError("invalid_pptx")
    partial.replace(dest)


def _is_zip(path: Path) -> bool:
    try:
        header = path.read_bytes()[:4]
    except OSError:
        return False
    return header.startswith(b"PK")


def _load_or_extract_slides(deck_dir: Path, pptx_path: Path) -> tuple[SlideText, ...]:
    content_path = deck_dir / "content.json"
    if content_path.is_file():
        try:
            return _read_content(content_path)
        except (OSError, json.JSONDecodeError, KeyError, TypeError, ValueError):
            content_path.unlink(missing_ok=True)
    slides = extract_slides(pptx_path)
    payload = {
        "version": _CACHE_VERSION,
        "slides": [
            {"number": slide.number, "title": slide.title, "body": slide.body, "notes": slide.notes}
            for slide in slides
        ],
    }
    _atomic_json(content_path, payload)
    return slides


def _read_content(path: Path) -> tuple[SlideText, ...]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if data.get("version") != _CACHE_VERSION:
        raise ValueError("stale cache")
    slides = []
    for item in data["slides"]:
        slides.append(
            SlideText(
                number=int(item["number"]),
                title=str(item["title"]),
                body=str(item["body"]),
                notes=str(item["notes"]),
            )
        )
    return tuple(slides)


def _load_cached_images(deck_dir: Path, slide_count: int) -> list[SlideImage] | None:
    manifest_path = deck_dir / "manifest.json"
    if not manifest_path.is_file():
        return None
    try:
        data = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None
    if data.get("version") != _CACHE_VERSION or data.get("text_only") is True:
        return None
    images: list[SlideImage] = []
    for item in data.get("images", []):
        path = deck_dir / item["file"]
        if not path.is_file():
            return None
        images.append(
            SlideImage(number=int(item["number"]), path=path, format=str(item["format"]))
        )
    paired, text_only = pair_images(slide_count, [(image.path, image.format) for image in images])
    if text_only:
        return None
    return paired


def _write_manifest(deck_dir: Path, images: list[SlideImage]) -> None:
    payload = {
        "version": _CACHE_VERSION,
        "text_only": False,
        "images": [
            {
                "number": image.number,
                "file": image.path.relative_to(deck_dir).as_posix(),
                "format": image.format,
            }
            for image in images
        ],
    }
    _atomic_json(deck_dir / "manifest.json", payload)


def _atomic_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    partial = path.with_suffix(path.suffix + ".part")
    partial.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    partial.replace(path)


def _convert_to_pdf(
    binary: str, pptx_path: Path, output_dir: Path, profile: Path, timeout: int
) -> None:
    command = [
        binary,
        "--headless",
        "--norestore",
        "--nolockcheck",
        "--nologo",
        f"-env:UserInstallation={_file_uri(profile)}",
        "--convert-to",
        "pdf",
        "--outdir",
        str(output_dir.resolve()),
        str(pptx_path.resolve()),
    ]
    env = os.environ.copy()
    env["HOME"] = str(profile)
    env["SAL_USE_VCLPLUGIN"] = "svp"
    process = subprocess.Popen(
        command,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        env=env,
        start_new_session=True,
    )
    try:
        process.communicate(timeout=timeout)
    except subprocess.TimeoutExpired as exc:
        _kill_process_group(process)
        raise RenderError("libreoffice_timeout") from exc
    if process.returncode != 0:
        raise RenderError("libreoffice_failed")


def _kill_process_group(process: subprocess.Popen) -> None:
    try:
        os.killpg(process.pid, signal.SIGKILL)
    except ProcessLookupError:
        return
    process.communicate()


def _find_pdf(output_dir: Path, pptx_path: Path) -> Path:
    expected = output_dir / f"{pptx_path.stem}.pdf"
    if expected.is_file():
        return expected
    pdfs = sorted(output_dir.glob("*.pdf"))
    if not pdfs:
        raise RenderError("libreoffice_failed")
    return pdfs[0]


def _pdf_to_images(pdf_path: Path, slide_dir: Path) -> list[tuple[Path, str]]:
    try:
        from pdf2image import convert_from_path
        from pdf2image.exceptions import PDFInfoNotInstalledError, PDFPageCountError
    except ImportError as exc:
        raise RenderError("pdf2image_not_installed") from exc
    slide_dir.mkdir(parents=True, exist_ok=True)
    try:
        pages = convert_from_path(str(pdf_path), dpi=110, fmt="png")
    except PDFInfoNotInstalledError as exc:
        raise RenderError("poppler_not_installed") from exc
    except PDFPageCountError as exc:
        raise RenderError("pdf_unreadable") from exc
    except Exception as exc:
        raise RenderError("pdf_render_failed") from exc
    rendered: list[tuple[Path, str]] = []
    for index, page in enumerate(pages, start=1):
        rendered.append(_save_image(page, slide_dir / f"slide-{index:03d}"))
    return rendered


def _save_image(image, dest_stem: Path) -> tuple[Path, str]:
    """Save a rendered page. Keep a PNG, and use JPEG when the PNG is large."""
    rgb = image.convert("RGB")
    rgb.thumbnail((1024, 1024))
    png_path = dest_stem.with_suffix(".png")
    rgb.save(png_path, format="PNG", optimize=True)
    if png_path.stat().st_size <= 450_000:
        return png_path, "png"
    jpeg_path = dest_stem.with_suffix(".jpg")
    png_path.unlink(missing_ok=True)
    quality = 75
    while quality >= 45:
        rgb.save(jpeg_path, format="JPEG", quality=quality, optimize=True)
        if jpeg_path.stat().st_size <= 450_000:
            return jpeg_path, "jpeg"
        quality -= 15
    smaller = rgb.copy()
    smaller.thumbnail((768, 768))
    smaller.save(jpeg_path, format="JPEG", quality=50, optimize=True)
    return jpeg_path, "jpeg"


def _slide_title(slide) -> tuple[str, int | None]:
    try:
        title_shape = slide.shapes.title
    except Exception:
        return "", None
    if title_shape is None:
        return "", None
    text = _clean(getattr(title_shape, "text", ""))
    return text, getattr(title_shape, "shape_id", None)


def _slide_notes(slide) -> str:
    try:
        if not slide.has_notes_slide:
            return ""
        text = _clean(slide.notes_slide.notes_text_frame.text)
    except Exception:
        return ""
    return text


def _chart_title(shape) -> str:
    try:
        chart = shape.chart
        if chart.has_title:
            return _clean(chart.chart_title.text_frame.text)
    except Exception:
        return ""
    return ""


def _clean(value: str | None) -> str:
    text = (value or "").strip()
    if text.lower().rstrip(".") in _PLACEHOLDERS:
        return ""
    return text


def _file_uri(path: Path) -> str:
    return "file://" + quote(path.resolve().as_posix(), safe="/")


def _lock_for(key: str) -> threading.Lock:
    with _URL_LOCKS_GUARD:
        lock = _URL_LOCKS.get(key)
        if lock is None:
            lock = threading.Lock()
            _URL_LOCKS[key] = lock
        return lock
